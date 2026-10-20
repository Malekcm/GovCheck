from pathlib import Path
import sys
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx_table_layout import add_fixed_table

root=Path(__file__).resolve().parents[1]
p=root/'templates'/'Federal_Opportunity_Research_Template.docx'
d=Document();sec=d.sections[0]
sec.left_margin=Inches(.75);sec.right_margin=Inches(.75);sec.top_margin=Inches(.75);sec.bottom_margin=Inches(.7)
for name in ['Normal','Body Text']:
 s=d.styles[name];s.font.name='Carlito';s.font.size=Pt(9.5);s.font.color.rgb=RGBColor(21,37,62);s.paragraph_format.space_after=Pt(5)
for name,size in [('Title',25),('Heading 1',15),('Heading 2',12)]:
 s=d.styles[name];s.font.name='Carlito';s.font.size=Pt(size);s.font.color.rgb=RGBColor(21,37,62);s.font.bold=True;s.paragraph_format.space_before=Pt(9);s.paragraph_format.space_after=Pt(6)
h=sec.header.paragraphs[0];h.alignment=WD_ALIGN_PARAGRAPH.RIGHT;h.add_run('FEDERAL OPPORTUNITY CAPTURE  |  INTERNAL WORKING TEMPLATE').font.size=Pt(7)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT;f.add_run('As-of date: [YYYY-MM-DD]    |    Page ').font.size=Pt(7)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)

def para(t):d.add_paragraph(t)
def heading(x):d.add_heading(x,level=1)
def sub(x):d.add_heading(x,level=2)
def tab(h,r,w):add_fixed_table(d,h,r,w);para('')

d.add_paragraph('GOVERNMENT BUSINESS DEVELOPMENT / CAPTURE INTELLIGENCE',style='Subtitle')
d.add_heading('[Agency / Opportunity Name]',0)
para('[Opportunity ID / SAM notice / forecast ID]   |   Research as-of [date]')
para('Evidence-backed working draft. Not a proposal. Forecast or predecessor terms are not evidence of current RFP requirements.')
heading('1. Leadership summary and recommended gate')
para('[Summary: pursuit rationale, fundamental blockers, current recommendation and next gate owner.]')
tab(['Issue','Current assessment','Evidence or next action'],[['Current RFP','OPEN','Confirm SAM or gated portal record, amend and deadline'],['Eligibility','OPEN','Exact prime vehicle/entity/domain/set-aside verified by contracts'],['Fit','OPEN','Approved workstream evidence and SME availability']], [1.65,1.4,3.95])
heading('2. Acquisition fingerprint and timeline')
tab(['Field','Finding','Source / caveat'],[['Buyer','OPEN','Agency notice'],['Solicitation and stage','OPEN','No unsupported claims'],['Vehicle / domain','OPEN','Verify eligible entity and contract'],['NAICS / size / set-aside','OPEN','Current official source'],['Value / contract type','OPEN','Separate forecast, ceiling, obligations and funding'],['Release, deadline, award, POP','OPEN','Current notice and amendments']], [1.7,2.8,2.5])
heading('3. Solicitation trace and government contacts')
tab(['Source/channel','Last checked / finding','Specific next action'],[['SAM.gov','OPEN','Check exact ID and aliases'],['Official agency / forecast','OPEN','Version / change log'],['Gated eBuy/portal','UNVERIFIED','Assigned contractor admin to check actual lane']], [1.7,2.45,2.85])
tab(['Contact / role','Verified channel','Engagement question / restriction'],[['Contracting officer','OPEN','Verify current name and authorized inquiry process'],['Program office','OPEN','Customer pain points and scope'],['Small business specialist','OPEN','Set-aside, engagement and eligibility']], [1.7,2.35,2.95])
heading('4. Incumbent, recompete and competitor research')
tab(['Evidence','Current finding','Source and caveat'],[['Predecessor contract','OPEN','Parent IDV and dates'],['Current incumbent','OPEN','Verified entity and prime/sub role'],['Obligations and ceiling','OPEN','Never conflate values'],['Major related awards','OPEN','Distinguish old and current competition']], [1.65,2.5,2.85])
heading('5. Scope and mission')
para('[Current PWS task families, service goals, deliverables, dependencies and historical-only task hypotheses. Cite every authoritative section.]')
heading('6. Company strengths and editable requirement matrix')
para('STRONG can reflect an internally confirmed company strength while approved past-performance artifacts are still outstanding. Do not call unapproved examples independently verified.')
tab(['Requirement family','Fit','Evidence / missing proof'],[['Program management','UNKNOWN','Contract, role, approved outcomes'],['Data and analytics','UNKNOWN','Relevant sources, accepted deliverables'],['Data / IT architecture','UNKNOWN','Formal designs, named technical leader'],['Specialized workstream','UNKNOWN','Qualified partner / SME plan']], [2.1,1.0,3.9])
sub('Granular traceability')
tab(['PWS/task','Company fit','Example / approved outcome','Required artifact/owner'],[['OPEN','UNKNOWN','OPEN','Evidence ID + assigned owner']], [1.6,.85,2.2,2.35])
heading('7. Qualification hard gates, economics and risk')
tab(['Mandatory gate','Status','Proof and approving owner'],[['Current RFP / due date','OPEN','Contracts'],['Vehicle / size / set aside','OPEN','Legal + Contracts'],['Security / OCI / workshare','OPEN','Legal + Security'],['Staffing / talent','OPEN','Delivery'],['FFP / cash / margin','OPEN','Finance + executive sponsor']], [2.05,1.0,3.95])
heading('8. BD questions with answers and recommendations')
tab(['ID','Question','Current answer','Evidence to request / owner'],[['A01','What is current solicitation status?','OPEN','Live notice + amendments / Contracts'],['F10','Which core strengths are approved?','OPEN','Past-performance package / Practice lead']], [.55,2.25,1.0,3.2])
heading('9. Immediate capture plan')
tab(['Priority','Request or task','Owner','Acceptance standard'],[['P0','Confirm live order and domain','Contracts','Dated screenshot or official posting'],['P0','Validate company capability references','Practice lead','Approved evidence cards'],['P1','Map incumbent / predecessor','Capture','Dated award and PWS citations']], [.7,2.4,1.15,2.75])
heading('10. Sources, open evidence and change log')
tab(['ID','Source or evidence','URL / file location and checked-on date'],[['S1','Official opportunity / agency notice','[Direct verified URL]'],['I1','Approved company profile','[Approval status / date]']], [.55,2.35,4.1])
para('Caveat: replace all OPEN placeholders with evidence or explicit missing requests. Render all Word pages and inspect table alignment before delivering.')
d.save(p)
print(p)
